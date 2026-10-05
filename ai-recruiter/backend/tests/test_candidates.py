import sys
import os
import io
import uuid
import pytest
import fitz
import docx

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from fastapi.testclient import TestClient
from app.main import app


def test_candidates_workflow():
    with TestClient(app) as client:
        # 1. Login as Recruiter
        login_res = client.post("/api/v1/auth/login", json={
            "email": "recruiter@smartrecruit.ai",
            "password": "Recruiter@123456"
        })
        assert login_res.status_code == 200
        token = login_res.json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}

        # 2. Upload TXT resume
        txt_content = (
            "Jonathan Archer\n"
            "jonathan.archer@starfleet.io | +1 555-014-9988\n"
            "Austin, TX\n\n"
            "SUMMARY: Senior Python and Cloud Backend Architect with 6+ years of experience.\n"
            "EDUCATION: Masters in Computer Science, Starfleet Academy\n"
            "SKILLS: Python, FastAPI, Django, PostgreSQL, Docker, AWS, Redis, Kubernetes\n"
        )
        txt_file = io.BytesIO(txt_content.encode("utf-8"))
        res_txt = client.post(
            "/api/v1/candidates/upload-resume",
            files={"file": ("jonathan_archer_resume.txt", txt_file, "text/plain")},
            headers=headers
        )
        assert res_txt.status_code == 201, res_txt.text
        txt_data = res_txt.json()
        cand = txt_data["candidate"]
        assert cand["name"] == "Jonathan Archer"
        assert cand["email"] == "jonathan.archer@starfleet.io"
        assert cand["total_experience"] == 6.0
        assert cand["education_level"] == "MASTERS"
        assert len(cand["skills"]) >= 5
        cand_id = cand["id"]

        # 3. Upload dynamic PDF resume (generated using PyMuPDF)
        pdf_doc = fitz.open()
        pdf_page = pdf_doc.new_page()
        pdf_text = (
            "Beverly Crusher\n"
            "beverly.crusher@enterprise.org | +1 415-555-2244\n\n"
            "Chief Data Analyst with 4.5 years of experience in healthcare informatics.\n"
            "Education: Bachelor of Science in Information Analytics\n"
            "Skills: SQL, Python, Tableau, Power BI, Statistical Analysis, Data Visualization\n"
        )
        pdf_page.insert_text((50, 72), pdf_text)
        pdf_bytes = pdf_doc.write()
        pdf_doc.close()

        res_pdf = client.post(
            "/api/v1/candidates/upload-resume",
            files={"file": ("beverly_crusher.pdf", io.BytesIO(pdf_bytes), "application/pdf")},
            headers=headers
        )
        assert res_pdf.status_code == 201, res_pdf.text
        pdf_data = res_pdf.json()["candidate"]
        assert pdf_data["name"] == "Beverly Crusher"
        assert pdf_data["email"] == "beverly.crusher@enterprise.org"
        assert pdf_data["total_experience"] == 4.5
        assert any(s["skill"] == "Tableau" for s in pdf_data["skills"])

        # 4. Upload dynamic DOCX resume (generated using python-docx)
        doc = docx.Document()
        doc.add_heading("Geordi La Forge", 0)
        doc.add_paragraph("geordi.laforge@enterprise.org | +1 512-555-7788 | Remote")
        doc.add_paragraph("DevOps & Cloud Engineer with 5+ years of experience in infrastructure automation.")
        doc.add_paragraph("Education: Master of Science in Systems Engineering")
        doc.add_paragraph("Core competencies: Kubernetes, Docker, AWS, Terraform, CI/CD, Linux, Prometheus")
        docx_buf = io.BytesIO()
        doc.save(docx_buf)
        docx_bytes = docx_buf.getvalue()

        res_docx = client.post(
            "/api/v1/candidates/upload-resume",
            files={"file": ("geordi_laforge.docx", io.BytesIO(docx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            headers=headers
        )
        assert res_docx.status_code == 201, res_docx.text
        docx_data = res_docx.json()["candidate"]
        assert docx_data["name"] == "Geordi La Forge"
        assert docx_data["email"] == "geordi.laforge@enterprise.org"
        assert any(s["skill"] == "Kubernetes" for s in docx_data["skills"])

        # 5. Invalid file type rejection
        bad_file = io.BytesIO(b"binary exe content")
        bad_res = client.post(
            "/api/v1/candidates/upload-resume",
            files={"file": ("malicious.exe", bad_file, "application/octet-stream")},
            headers=headers
        )
        assert bad_res.status_code == 400

        # 6. List and filter candidates
        list_res = client.get("/api/v1/candidates?search=Starfleet", headers=headers)
        assert list_res.status_code == 200
        assert len(list_res.json()) >= 1

        skill_filter_res = client.get("/api/v1/candidates?skill=Kubernetes", headers=headers)
        assert skill_filter_res.status_code == 200
        assert len(skill_filter_res.json()) >= 1

        # 7. Seed 20 realistic candidate profiles
        seed_res = client.post("/api/v1/candidates/seed", headers=headers)
        assert seed_res.status_code == 200
        assert len(seed_res.json()) >= 20

        # 8. Update candidate profile
        update_res = client.put(
            f"/api/v1/candidates/{cand_id}",
            json={"location": "Austin, TX (Hybrid)", "total_experience": 7.0},
            headers=headers
        )
        assert update_res.status_code == 200
        assert update_res.json()["location"] == "Austin, TX (Hybrid)"
        assert update_res.json()["total_experience"] == 7.0

        # 9. Delete candidate
        del_res = client.delete(f"/api/v1/candidates/{cand_id}", headers=headers)
        assert del_res.status_code == 200
        check_del = client.get(f"/api/v1/candidates/{cand_id}", headers=headers)
        assert check_del.status_code == 404
